from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

from slides.utils import SlideText

IMAGE_PATH = Path(__file__).resolve().parent.parent / "assets" / "bunnies.jpg"

BULLETS = [
    "Start with one pair of newborn rabbits",
    "Each month, every mature pair produces one new pair",
    "Rabbits mature after one month and never die",
    "How many pairs after 12 months? 1, 1, 2, 3, 5, 8, …",
]


def build(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[5])  # Title Only layout
    slide.shapes.title.text = "The Rabbit Problem (Liber Abaci, 1202)"

    left = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(4.8), Inches(4.5))
    frame = left.text_frame
    frame.word_wrap = True
    for i, bullet in enumerate(BULLETS):
        paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        paragraph.text = bullet
        paragraph.level = 0
        paragraph.font.size = Pt(20)
        SlideText.add_bullet(paragraph)
        if i > 0:
            paragraph.space_before = Pt(12)

    # Photo is wide (1280 × 634); width only, height scales proportionally.
    slide.shapes.add_picture(str(IMAGE_PATH), Inches(5.8), Inches(1.8), width=Inches(3.7))

    caption = slide.shapes.add_textbox(Inches(5.8), Inches(6.3), Inches(3.7), Inches(0.5))
    caption.text_frame.word_wrap = True
    run = caption.text_frame.paragraphs[0]
    run.text = "Photo: James Johnstone, CC BY 2.0, via Wikimedia Commons"
    run.font.size = Pt(10)
