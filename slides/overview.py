"""Finale slide: every algorithm at a glance.

A roadmap diagram (generated with PIL at build time into target/) compresses
the deck's four implementations into one visual: the complexity ladder from
exponential to constant.
"""

from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches, Pt

from slides.utils import SlideText, draw_fraction, gif_font

TARGET_DIR = Path(__file__).resolve().parent.parent / "target"
ROADMAP_PATH = TARGET_DIR / "roadmap.png"

IMG_W, IMG_H = 1200, 520
PANEL_W, GAP, X0 = 280, 20, 10

INK, GREY = (31, 31, 31), (110, 110, 110)
PANEL_FILL, PANEL_LINE = (245, 247, 250), (180, 185, 190)
BLUE_FILL, BLUE_LINE = (219, 234, 249), (31, 78, 121)
GREEN_FILL, GREEN_LINE = (198, 224, 180), (56, 118, 29)
ORANGE_FILL = (255, 217, 102)

BULLETS = [
    "Exponential → linear → constant: each step removes repeated work",
    "First pay memory for speed (the table), then shrink the memory (two variables)",
    "Binet wins on paper — floating point rounds it off past n ≈ 70",
]


def _panel(draw: ImageDraw.ImageDraw, i: int, title: str) -> float:
    x0 = X0 + i * (PANEL_W + GAP)
    draw.rectangle([x0, 10, x0 + PANEL_W, IMG_H - 10], fill=PANEL_FILL, outline=PANEL_LINE, width=2)
    draw.text((x0 + PANEL_W / 2, 44), title, font=gif_font(28), fill=INK, anchor="ma")
    return x0


def _complexity(draw: ImageDraw.ImageDraw, x0: float, text: str) -> None:
    draw.text((x0 + PANEL_W / 2, 442), text, font=gif_font(23), fill=INK, anchor="ma")


def _node(draw: ImageDraw.ImageDraw, x: float, y: float, n: int, r: int = 24) -> None:
    draw.ellipse([x - r, y - r, x + r, y + r], fill="white", outline=BLUE_LINE, width=2)
    draw.text((x, y), str(n), font=gif_font(18), fill=INK, anchor="mm")


def _tree(draw: ImageDraw.ImageDraw, x0: float) -> None:
    cx = x0 + PANEL_W / 2
    nodes = [  # (label, x, y); values repeat, so no dict keyed by value
        (5, cx, 125),
        (4, cx - 60, 210),
        (3, cx + 60, 210),
        (3, cx - 92, 290),
        (2, cx - 28, 290),
        (2, cx + 28, 290),
        (1, cx + 92, 290),
    ]
    for a, b in ((0, 1), (0, 2), (1, 3), (1, 4), (2, 5), (2, 6)):
        _, ax, ay = nodes[a]
        _, bx, by = nodes[b]
        draw.line([(ax, ay + 24), (bx, by - 24)], fill=GREY, width=2)
    for n, x, y in nodes:
        _node(draw, x, y, n)
    draw.text((cx, 348), "...", font=gif_font(24), fill=GREY, anchor="ma")


def _boxes(draw: ImageDraw.ImageDraw, x0: float, values: list, window: tuple[int, ...] = ()) -> None:
    n, bw, gap = len(values), 52 if len(values) <= 4 else 36, 6
    total = n * (bw + gap) - gap
    x = x0 + (PANEL_W - total) / 2
    y, h = 190, 64
    for j, v in enumerate(values):
        fill = ORANGE_FILL if j in window else GREEN_FILL
        draw.rectangle([x, y, x + bw, y + h], fill=fill, outline=GREEN_LINE if fill == GREEN_FILL else BLUE_LINE, width=2)
        draw.text((x + bw / 2, y + h / 2), str(v), font=gif_font(22), fill=INK, anchor="mm")
        if j in window:
            draw.text((x + bw / 2, y + h + 22), "ab"[window.index(j)], font=gif_font(18), fill=INK, anchor="ma")
        x += bw + gap


def _binet(draw: ImageDraw.ImageDraw, x0: float) -> None:
    font = gif_font(26)
    num, den = "(φⁿ − ψⁿ)", "√5"
    nb = draw.textbbox((0, 0), num, font=font)
    db = draw.textbbox((0, 0), den, font=font)
    w = max(nb[2] - nb[0], db[2] - db[0])
    draw_fraction(draw, x0 + (PANEL_W - w) / 2, 150, num, den, font, rule=3)
    draw.text((x0 + PANEL_W / 2, 320), "just plug in n", font=gif_font(20), fill=GREY, anchor="ma")


def _make_roadmap() -> None:
    img = Image.new("RGB", (IMG_W, IMG_H), "white")
    draw = ImageDraw.Draw(img)
    x0 = _panel(draw, 0, "Naive recursion")
    _tree(draw, x0)
    _complexity(draw, x0, "O(2ⁿ) time")
    x0 = _panel(draw, 1, "Memoization")
    _boxes(draw, x0, [0, 1, 1, 2, 3, 5])
    _complexity(draw, x0, "O(n) time · O(n) space")
    x0 = _panel(draw, 2, "Two variables")
    _boxes(draw, x0, [2, 3, 5, 8], window=(2, 3))
    _complexity(draw, x0, "O(n) time · O(1) space")
    x0 = _panel(draw, 3, "Binet")
    _binet(draw, x0)
    _complexity(draw, x0, "O(1) time")
    for gx in (X0 + PANEL_W + GAP / 2 + k * (PANEL_W + GAP) for k in range(3)):
        draw.text((gx, 240), "→", font=gif_font(34), fill=GREY, anchor="mm")
    img.save(ROADMAP_PATH)


def build(prs: Presentation) -> None:
    TARGET_DIR.mkdir(parents=True, exist_ok=True)
    _make_roadmap()

    slide = prs.slides.add_slide(prs.slide_layouts[5])  # Title Only layout
    slide.shapes.title.text = "Every Algorithm at a Glance"

    # Roadmap is 1200×520; width only, height scales proportionally.
    slide.shapes.add_picture(str(ROADMAP_PATH), Inches(0.5), Inches(1.4), width=Inches(9.0))

    box = slide.shapes.add_textbox(Inches(0.5), Inches(5.45), Inches(9.0), Inches(1.85))
    frame = box.text_frame
    frame.word_wrap = True
    for i, bullet in enumerate(BULLETS):
        paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        paragraph.text = bullet
        paragraph.level = 0
        paragraph.font.size = Pt(18)
        SlideText.add_bullet(paragraph)
        if i > 0:
            paragraph.space_before = Pt(6)
