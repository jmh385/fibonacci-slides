from pptx import Presentation
from pptx.util import Inches, Pt

from slides.utils import SlideText

BULLETS = [
    "Each term is the sum of the two preceding terms: F(n) = F(n − 1) + F(n − 2)",
    "Base cases anchor the sequence: F(0) = 0, F(1) = 1",
    "Rabbit pairs month by month: 1, 1, 2, 3, 5, 8, 13, …",
]

N_VALUES = list(range(0, 11))
F_VALUES = [0, 1, 1, 2, 3, 5, 8, 13, 21, 34, 55]


def build(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[5])  # Title Only layout
    slide.shapes.title.text = "The Recurrence: F(n) = F(n − 1) + F(n − 2)"

    box = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(9.0), Inches(2.3))
    frame = box.text_frame
    frame.word_wrap = True
    for i, bullet in enumerate(BULLETS):
        paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        paragraph.text = bullet
        paragraph.level = 0
        paragraph.font.size = Pt(20)
        SlideText.add_bullet(paragraph)
        if i > 0:
            paragraph.space_before = Pt(12)

    rows, cols = 2, len(N_VALUES) + 1
    table_shape = slide.shapes.add_table(rows, cols, Inches(0.5), Inches(4.4), Inches(9.0), Inches(1.4))
    table = table_shape.table
    table.cell(0, 0).text = "n"
    table.cell(1, 0).text = "F(n)"
    for j, (n, f) in enumerate(zip(N_VALUES, F_VALUES)):
        table.cell(0, j + 1).text = str(n)
        table.cell(1, j + 1).text = str(f)
    for row in range(rows):
        for col in range(cols):
            cell = table.cell(row, col)
            for paragraph in cell.text_frame.paragraphs:
                paragraph.alignment = 2  # center
                for run in paragraph.runs:
                    run.font.size = Pt(18)
                    if row == 0 or col == 0:
                        run.font.bold = True
