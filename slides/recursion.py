"""Naive recursion slide.

python-pptx cannot author PowerPoint entrance animations, so the "animation"
is an animated GIF: the fib(5) call tree is drawn with PIL frame by frame
(calls expand top-down, results bubble bottom-up) and embedded with
add_picture. It plays automatically in slideshow mode.
The GIF is regenerated into target/ on every build, so it is never stale.
"""

from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from slides.utils import SlideText, gif_font

TARGET_DIR = Path(__file__).resolve().parent.parent / "target"
GIF_PATH = TARGET_DIR / "recursion-tree.gif"

PSEUDOCODE = [
    "fib(n):",
    "  if n <= 1: return n",
    "  return fib(n-1) + fib(n-2)",
]

BULLETS = [
    "Top-down: F(n) splits into F(n−1) + F(n−2)",
    "Same subproblems recomputed — fib(3) runs twice for fib(5)",
    "Time blows up ~O(2ⁿ): needs memoization →",
]

# GIF layout
CALL_N = 5
IMG_W, IMG_H = 800, 600
NODE_R = 34
TOP, Y_STEP = 60, 105
STATUS_Y = 545

BLUE_FILL, BLUE_LINE = (219, 234, 249), (31, 78, 121)
ORANGE_FILL = (255, 217, 102)
GREEN_FILL, GREEN_LINE = (198, 224, 180), (56, 118, 29)
INK = (31, 31, 31)


def _fib(n: int) -> int:
    a, b = 0, 1
    for _ in range(n):
        a, b = b, a + b
    return a


class _Node:
    def __init__(self, n: int, depth: int) -> None:
        self.n = n
        self.depth = depth
        self.children = [_Node(n - 1, depth + 1), _Node(n - 2, depth + 1)] if n >= 2 else []
        self.x = 0.0

    def preorder(self) -> list["_Node"]:
        return [self] + [m for c in self.children for m in c.preorder()]

    def postorder(self) -> list["_Node"]:
        return [m for c in self.children for m in c.postorder()] + [self]


def _layout(root: _Node) -> None:
    leaves = [m for m in root.preorder() if not m.children]
    step = IMG_W / len(leaves)
    for i, leaf in enumerate(leaves):
        leaf.x = step * (i + 0.5)

    def place(node: _Node) -> None:
        for child in node.children:
            place(child)
        if node.children:
            node.x = sum(c.x for c in node.children) / len(node.children)

    place(root)


def _draw_frame(shown: set[int], valued: set[int], status: str, nodes: list[_Node]) -> Image.Image:
    img = Image.new("RGB", (IMG_W, IMG_H), "white")
    draw = ImageDraw.Draw(img)
    index = {id(m): i for i, m in enumerate(nodes)}
    for i, node in enumerate(nodes):
        if i not in shown:
            continue
        y = TOP + node.depth * Y_STEP
        for child in node.children:
            j = index[id(child)]
            if j in shown:
                cy = TOP + child.depth * Y_STEP
                draw.line([(node.x, y + NODE_R), (child.x, cy - NODE_R)], fill=(128, 128, 128), width=2)
    label = gif_font(19)
    for i, node in enumerate(nodes):
        if i not in shown:
            continue
        y = TOP + node.depth * Y_STEP
        if i in valued:
            fill, line, lines = GREEN_FILL, GREEN_LINE, [f"fib({node.n})", f"= {_fib(node.n)}"]
        elif i == max(shown):
            fill, line, lines = ORANGE_FILL, BLUE_LINE, [f"fib({node.n})"]
        else:
            fill, line, lines = BLUE_FILL, BLUE_LINE, [f"fib({node.n})"]
        draw.ellipse([node.x - NODE_R, y - NODE_R, node.x + NODE_R, y + NODE_R], fill=fill, outline=line, width=2)
        heights = [draw.textbbox((0, 0), t, font=label)[3] for t in lines]
        ty = y - sum(heights) / 2 - (4 if len(lines) > 1 else 0)
        for t, h in zip(lines, heights):
            draw.text((node.x, ty), t, font=label, fill=INK, anchor="ma")
            ty += h + 4
    status_font = gif_font(21)
    draw.text((IMG_W / 2, STATUS_Y), status, font=status_font, fill=INK, anchor="ma")
    return img


def _make_gif() -> None:
    root = _Node(CALL_N, 0)
    _layout(root)
    pre = root.preorder()
    post = root.postorder()
    frames: list[Image.Image] = []
    durations: list[int] = []
    for k in range(1, len(pre) + 1):
        node = pre[k - 1]
        frames.append(_draw_frame(set(range(k)), set(), f"call {k}/{len(pre)}: fib({node.n})", pre))
        durations.append(600)
    by_depth: dict[int, list[_Node]] = {}
    for node in post:
        by_depth.setdefault(node.depth, []).append(node)
    valued: set[int] = set()
    index = {id(m): i for i, m in enumerate(pre)}
    for depth in sorted(by_depth, reverse=True):
        valued.update(index[id(m)] for m in by_depth[depth])
        frames.append(_draw_frame(set(range(len(pre))), set(valued), "results bubble up…", pre))
        durations.append(800)
    # Single hold frame: Pillow merges pixel-identical consecutive frames,
    # so two identical holds would collapse into one.
    frames.append(_draw_frame(set(range(len(pre))), set(valued), f"fib({CALL_N}) = {_fib(CALL_N)}", pre))
    durations.append(5000)
    frames[0].save(GIF_PATH, save_all=True, append_images=frames[1:], duration=durations, loop=0)


def build(prs: Presentation) -> None:
    TARGET_DIR.mkdir(parents=True, exist_ok=True)
    _make_gif()

    slide = prs.slides.add_slide(prs.slide_layouts[5])  # Title Only layout
    slide.shapes.title.text = "Naive Recursion: fib(n) = fib(n−1) + fib(n−2)"

    code = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(4.2), Inches(1.6))
    code.fill.solid()
    code.fill.fore_color.rgb = RGBColor(243, 244, 246)
    frame = code.text_frame
    frame.word_wrap = True
    for i, line in enumerate(PSEUDOCODE):
        paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        paragraph.text = line
        run = paragraph.runs[0]
        run.font.size = Pt(18)
        run.font.name = "Consolas"

    box = slide.shapes.add_textbox(Inches(0.5), Inches(3.4), Inches(4.2), Inches(2.8))
    frame = box.text_frame
    frame.word_wrap = True
    for i, bullet in enumerate(BULLETS):
        paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        paragraph.text = bullet
        paragraph.level = 0
        paragraph.font.size = Pt(20)
        SlideText.add_bullet(paragraph)
        if i > 0:
            paragraph.space_before = Pt(12)

    # GIF is 800×600; width only, height scales proportionally.
    slide.shapes.add_picture(str(GIF_PATH), Inches(5.0), Inches(1.5), width=Inches(4.5))

    caption = slide.shapes.add_textbox(Inches(5.0), Inches(5.1), Inches(4.5), Inches(0.8))
    caption.text_frame.word_wrap = True
    run = caption.text_frame.paragraphs[0]
    run.text = "Call tree for fib(5), call by call (animated GIF — plays in slideshow mode)"
    run.font.size = Pt(10)
