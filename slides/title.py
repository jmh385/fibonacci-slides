from pptx import Presentation


def build(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[0])  # Title Slide layout
    slide.shapes.title.text = "Fibonacci Algorithms"
    slide.placeholders[1].text = "Comparing implementations and visualizing them"
