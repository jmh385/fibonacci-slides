"""Binet's formula slide: closed form, no loop, no recursion.

python-pptx cannot author real equations, so the formula is rendered with PIL
(full division lines) into target/ at build time and embedded as an image.
"""

from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches, Pt

from slides.utils import SlideText, draw_fraction, gif_font

IMAGE_PATH = Path(__file__).resolve().parent.parent / "assets" / "binet.jpg"
TARGET_DIR = Path(__file__).resolve().parent.parent / "target"
FORMULA_PATH = TARGET_DIR / "binet-formula.png"

IMG_W, IMG_H = 1000, 360
INK = (31, 31, 31)

BULLETS = [
    "No loop, no recursion — just plug in n",
    "Named for Jacques Binet (1786–1856); de Moivre knew it a century earlier",
]


def _width(draw: ImageDraw.ImageDraw, text: str, font) -> float:
    box = draw.textbbox((0, 0), text, font=font)
    return box[2] - box[0]


def _height(draw: ImageDraw.ImageDraw, text: str, font) -> float:
    box = draw.textbbox((0, 0), text, font=font)
    return box[3] - box[1]


def _make_formula() -> None:
    img = Image.new("RGB", (IMG_W, IMG_H), "white")
    draw = ImageDraw.Draw(img)

    main_font = gif_font(50)
    num, den, prefix = "(φⁿ − ψⁿ)", "√5", "F(n) = "
    frac_w = max(_width(draw, num, main_font), _width(draw, den, main_font))
    row_w = _width(draw, prefix, main_font) + 30 + frac_w
    x = (IMG_W - row_w) / 2
    y_top = 28
    _, frac_h = draw_fraction(draw, x + _width(draw, prefix, main_font) + 30, y_top, num, den, main_font)
    vinculum_y = y_top + _height(draw, num, main_font) + 14
    draw.text((x, vinculum_y), prefix, font=main_font, fill=INK, anchor="lm")

    small_font = gif_font(32)
    defs = [
        ("φ = ", "1 + √5", "2", "≈ 1.618"),
        ("ψ = ", "1 − √5", "2", "≈ −0.618"),
    ]
    blocks = []
    for sym, fnum, fden, val in defs:
        fw = max(_width(draw, fnum, small_font), _width(draw, fden, small_font))
        blocks.append((sym, fnum, fden, val, _width(draw, sym, small_font) + 16 + fw + 16 + _width(draw, val, small_font)))
    y_top = 205
    x = (IMG_W - (blocks[0][4] + 90 + blocks[1][4])) / 2
    for sym, fnum, fden, val, total in blocks:
        draw.text((x, y_top + 30), sym, font=small_font, fill=INK, anchor="lm")
        x += _width(draw, sym, small_font) + 16
        fw, _ = draw_fraction(draw, x, y_top, fnum, fden, small_font, pad=10, rule=3)
        x += fw + 16
        draw.text((x, y_top + 30), val, font=small_font, fill=INK, anchor="lm")
        x += _width(draw, val, small_font) + 90
    img.save(FORMULA_PATH)


def build(prs: Presentation) -> None:
    TARGET_DIR.mkdir(parents=True, exist_ok=True)
    _make_formula()

    slide = prs.slides.add_slide(prs.slide_layouts[5])  # Title Only layout
    slide.shapes.title.text = "Closed Form: Binet's Formula"

    # Portrait is small (161 × 227); keep it modest on the left.
    slide.shapes.add_picture(str(IMAGE_PATH), Inches(0.5), Inches(1.5), width=Inches(2.5))

    caption = slide.shapes.add_textbox(Inches(0.5), Inches(5.3), Inches(2.5), Inches(0.8))
    caption.text_frame.word_wrap = True
    run = caption.text_frame.paragraphs[0]
    run.text = "Portrait: unknown artist, 19th century (public domain, via Wikimedia Commons)"
    run.font.size = Pt(10)

    # Formula image is 1000×360; width only, height scales proportionally.
    slide.shapes.add_picture(str(FORMULA_PATH), Inches(3.4), Inches(1.5), width=Inches(6.1))

    box = slide.shapes.add_textbox(Inches(3.4), Inches(3.9), Inches(6.1), Inches(2.6))
    frame = box.text_frame
    frame.word_wrap = True
    for i, bullet in enumerate(BULLETS):
        paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        paragraph.text = bullet
        paragraph.level = 0
        paragraph.font.size = Pt(20)
        SlideText.add_bullet(paragraph)
        if i > 0:
            paragraph.space_before = Pt(12)
